import os
import subprocess
import sys

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install('python-docx')
    import docx

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def add_code_to_doc(doc, filename, filepath):
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
    except Exception as e:
        content = f"Error reading file: {e}"
    
    # Add title
    p = doc.add_paragraph()
    runner = p.add_run(filename)
    runner.bold = True
    runner.font.size = Pt(14)
    runner.font.color.rgb = RGBColor(0, 0, 128)
    
    # Add code
    p_code = doc.add_paragraph()
    run_code = p_code.add_run(content)
    run_code.font.name = 'Courier New'
    run_code.font.size = Pt(9)
    # add a page break after each file
    doc.add_page_break()

def main():
    doc = Document()
    head = doc.add_paragraph()
    head.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run_head = head.add_run("Academic Workload System Codebase")
    run_head.bold = True
    run_head.font.size = Pt(18)
    doc.add_page_break()

    base_dir = r"C:\Users\Sai Ishitha\.gemini\antigravity\scratch\academic_workload_system"
    
    files_to_include = [
        "app.py",
        "database.py",
        "ui_components.py",
        "templates/base.html",
        "templates/login.html",
        "templates/register.html",
        "templates/dashboard.html",
        "templates/tasks.html",
        "templates/exam.html",
    ]
    
    # We will also add static files if they exist
    for root, dirs, files in os.walk(os.path.join(base_dir, "static")):
        for file in files:
            if file.endswith('.css') or file.endswith('.js'):
                rel_path = os.path.relpath(os.path.join(root, file), base_dir)
                files_to_include.append(rel_path.replace('\\', '/'))

    for rel_path in files_to_include:
        filepath = os.path.join(base_dir, rel_path)
        if os.path.exists(filepath):
            add_code_to_doc(doc, rel_path, filepath)
            
    doc.save(os.path.join(base_dir, "Project_Code.docx"))
    print("Project_Code.docx generated successfully.")

if __name__ == "__main__":
    main()
