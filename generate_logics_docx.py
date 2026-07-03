import os
import subprocess
import sys

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install('python-docx')
    import docx

from docx import Document
from docx.shared import Pt, RGBColor

def main():
    doc = Document()
    
    # Title
    doc.add_heading('Smart Academic Workload & Exam Preparation System - Code Logics', 0)
    
    doc.add_paragraph('This document outlines the core business logic, database structures, and algorithms utilized in the Smart Academic Workload System.')
    
    # Section 1
    doc.add_heading('1. System Architecture', level=1)
    p = doc.add_paragraph()
    p.add_run('The application is built as a traditional web application combined with REST-like API endpoints for dynamic frontend data rendering.\n').bold = True
    p.add_run('• Framework: Flask (Python)\n')
    p.add_run('• Database: SQLite3 with pandas for querying data into DataFrames.\n')
    p.add_run('• Frontend Stack: HTML/CSS (Jinja2 Templates) + JavaScript (Chart.js / Vanilla JS).')

    # Section 2
    doc.add_heading('2. Database Models (database.py)', level=1)
    p2 = doc.add_paragraph('The system consists of three primary entities linked through foreign keys:\n')
    p2.add_run('• Users Table: ').bold = True
    p2.add_run('id, username, password (hashed), email, age, dob, gender, field_of_study.\n')
    p2.add_run('• Tasks Table: ').bold = True
    p2.add_run('id, user_id (FK), title, type, deadline, status, subject, estimated_hours.\n')
    p2.add_run('• Subjects Table: ').bold = True
    p2.add_run('id, user_id (FK), name, target_score, preparation_progress.')

    # Section 3
    doc.add_heading('3. Core Algorithms and Logics (app.py)', level=1)
    
    doc.add_heading('3.1 User Authentication', level=2)
    p3 = doc.add_paragraph()
    p3.add_run('Registration: ').bold = True
    p3.add_run('When a new user registers (/register), their plaintext password is obfuscated using werkzeug.security.generate_password_hash.\n')
    p3.add_run('Login/Session: ').bold = True
    p3.add_run('During login (/login), check_password_hash verifies the passwords. Upon success, the session is populated with the corresponding user_id.')

    doc.add_heading('3.2 Automated Stress Score Calculation', level=2)
    p4 = doc.add_paragraph('The /api/dashboard_data endpoint calculates a Stress Score (0 to 100) dynamically using an algorithm based on pending tasks, deadlines, and exam readiness:\n')
    p4.add_run('1. Volume of Work: Add 5 points for every pending task (maximum 40 points).\n')
    p4.add_run('2. Urgency Factor: Find tasks due within the next 3 days. Add 10 points for each urgent task (maximum 30 points).\n')
    p4.add_run('3. Exam Anxiety Factor: Calculates average exam readiness. If below 50%, the score increases proportionally by (50 - average) * 0.6.\n')
    p4.add_run('Classification:\n')
    p4.add_run('- Low Stress (< 30)\n- Moderate Stress (30 - 69)\n- High Stress (>= 70)')

    doc.add_heading('3.3 Workload Distribution', level=2)
    doc.add_paragraph('Using pandas, the system filters for Pending tasks and groups them by subject. It calculates the sum of estimated_hours for each subject. This aggregation is serialized into JSON to drive the charts.')

    doc.add_heading('3.4 Intelligent Reminders Generator', level=2)
    p5 = doc.add_paragraph('The /api/reminders endpoint calculates notifications based on:\n')
    p5.add_run('• Overdue Tasks: ').bold = True
    p5.add_run('Deadline has passed.\n')
    p5.add_run('• Urgent Tasks: ').bold = True
    p5.add_run('Due exactly today.\n')
    p5.add_run('• Upcoming Tasks: ').bold = True
    p5.add_run('Due in 1 to 3 days.\n')
    p5.add_run('• Off-Track / Low Readiness Exams: ').bold = True
    p5.add_run('If current preparation progress is far below the target score or critically below 50%.')

    out_file = r'C:\Users\Sai Ishitha\.gemini\antigravity\scratch\academic_workload_system\Code_Logics.docx'
    doc.save(out_file)
    print(f"Document saved successfully to {out_file}")

if __name__ == '__main__':
    main()
